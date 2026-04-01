#!/usr/bin/env python3
"""
Board Game Manual Generator
Genera manuales de juegos de mesa en formato .md y .docx
a partir de un archivo JSON con los datos del juego.

Uso:
    python generate_manual.py <input_json> [--output-dir <dir>] [--format md|docx|both]
"""

import json
import sys
import os
import re
import argparse
from datetime import datetime
from pathlib import Path


def load_game_data(json_path: str) -> dict:
    """Carga datos del juego desde un archivo JSON."""
    with open(json_path, "r", encoding="utf-8") as f:
        return json.load(f)


def build_components_table(components: list) -> str:
    """Construye una tabla Markdown de componentes."""
    if not components:
        return "_[COMPONENTES — pendiente de completar]_"
    lines = ["| Componente | Cantidad | Descripción |", "|---|---|---|"]
    for c in components:
        nombre = c.get("nombre", "")
        cantidad = c.get("cantidad", "")
        descripcion = c.get("descripcion", "")
        lines.append(f"| {nombre} | {cantidad} | {descripcion} |")
    return "\n".join(lines)


def build_setup_steps(steps: list) -> str:
    """Construye una lista numerada de pasos de setup."""
    if not steps:
        return "_[PASOS_SETUP — pendiente de completar]_"
    return "\n".join(f"{i+1}. {step}" for i, step in enumerate(steps))


def build_actions_sections(actions: list) -> str:
    """Construye secciones H3 para cada acción."""
    if not actions:
        return "_[ACCIONES — pendiente de completar]_"
    sections = []
    for action in actions:
        nombre = action.get("nombre", "")
        descripcion = action.get("descripcion", "")
        sections.append(f"### {nombre}\n\n{descripcion}")
    return "\n\n".join(sections)


def build_glossary_table(glossary: list) -> str:
    """Construye una tabla Markdown del glosario."""
    if not glossary:
        return "_[GLOSARIO — pendiente de completar]_"
    lines = ["| Término | Definición |", "|---|---|"]
    for item in glossary:
        termino = item.get("termino", "")
        definicion = item.get("definicion", "")
        lines.append(f"| {termino} | {definicion} |")
    return "\n".join(lines)


def build_faq_section(faq: list) -> str:
    """Construye la sección de preguntas frecuentes."""
    if not faq:
        return "_[FAQ — pendiente de completar]_"
    items = []
    for item in faq:
        pregunta = item.get("pregunta", "")
        respuesta = item.get("respuesta", "")
        items.append(f"**P: {pregunta}**\nR: {respuesta}")
    return "\n\n".join(items)


def preprocess_data(data: dict) -> dict:
    """
    Convierte campos estructurados (arrays de objetos) en texto Markdown
    para ser insertados en el template.
    """
    processed = dict(data)

    # Componentes
    if "COMPONENTES" in data and isinstance(data["COMPONENTES"], list):
        processed["TABLA_COMPONENTES"] = build_components_table(data["COMPONENTES"])
    else:
        processed["TABLA_COMPONENTES"] = "_[COMPONENTES — pendiente de completar]_"

    # Pasos de setup
    if "PASOS_SETUP" in data and isinstance(data["PASOS_SETUP"], list):
        processed["LISTA_PASOS_SETUP"] = build_setup_steps(data["PASOS_SETUP"])
    else:
        processed["LISTA_PASOS_SETUP"] = "_[PASOS_SETUP — pendiente de completar]_"

    # Acciones
    if "ACCIONES" in data and isinstance(data["ACCIONES"], list):
        processed["SECCIONES_ACCIONES"] = build_actions_sections(data["ACCIONES"])
    else:
        processed["SECCIONES_ACCIONES"] = "_[ACCIONES — pendiente de completar]_"

    # Glosario
    if "GLOSARIO" in data and isinstance(data["GLOSARIO"], list):
        processed["TABLA_GLOSARIO"] = build_glossary_table(data["GLOSARIO"])
    else:
        processed["TABLA_GLOSARIO"] = "_[GLOSARIO — pendiente de completar]_"

    # FAQ
    if "FAQ" in data and isinstance(data["FAQ"], list):
        processed["CONTENIDO_FAQ"] = build_faq_section(data["FAQ"])
    else:
        processed["CONTENIDO_FAQ"] = "_[FAQ — pendiente de completar]_"

    return processed


def render_template(template_path: str, data: dict) -> str:
    """Reemplaza placeholders {{KEY}} en el template con los valores del diccionario."""
    with open(template_path, "r", encoding="utf-8") as f:
        content = f.read()

    # Preprocesar datos estructurados
    processed = preprocess_data(data)

    for key, value in processed.items():
        placeholder = f"{{{{{key}}}}}"
        if isinstance(value, list):
            content = content.replace(placeholder, "\n".join(str(v) for v in value))
        elif value is None or value == "":
            content = content.replace(placeholder, f"_[{key} — pendiente de completar]_")
        else:
            content = content.replace(placeholder, str(value))

    # Marcar placeholders no reemplazados como pendientes
    remaining = re.findall(r"\{\{[A-Z_0-9]+\}\}", content)
    for placeholder in set(remaining):
        key_name = placeholder.strip("{}")
        content = content.replace(placeholder, f"_[{key_name} — pendiente de completar]_")

    return content


def save_markdown(content: str, output_path: str) -> None:
    """Guarda el manual como archivo Markdown."""
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"[OK] Manual Markdown guardado en: {output_path}")


def save_docx(content: str, output_path: str) -> None:
    """Convierte el contenido Markdown a .docx usando python-docx."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # Estilos base
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Calibri"
        font.size = Pt(11)

        lines = content.split("\n")
        i = 0
        while i < len(lines):
            line = lines[i]

            # Título principal (H1)
            if re.match(r"^# [^#]", line):
                heading = doc.add_heading(line[2:].strip(), level=1)
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # H2
            elif re.match(r"^## [^#]", line):
                doc.add_heading(line[3:].strip(), level=2)

            # H3
            elif re.match(r"^### [^#]", line):
                doc.add_heading(line[4:].strip(), level=3)

            # Blockquote (> texto)
            elif line.startswith("> "):
                para = doc.add_paragraph()
                run = para.add_run(line[2:].strip())
                run.italic = True

            # Separador horizontal
            elif line.strip() == "---":
                doc.add_paragraph("─" * 60)

            # Tabla Markdown (detectar por | en línea y separador en siguiente)
            elif "|" in line and i + 1 < len(lines) and re.match(r"^\|[-| ]+\|$", lines[i + 1].strip()):
                headers = [h.strip() for h in line.split("|") if h.strip()]
                i += 2  # Saltar línea de separador
                rows = []
                while i < len(lines) and "|" in lines[i]:
                    row = [c.strip() for c in lines[i].split("|") if c.strip()]
                    if row:
                        rows.append(row)
                    i += 1

                if headers:
                    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
                    table.style = "Table Grid"
                    hdr_cells = table.rows[0].cells
                    for j, header in enumerate(headers):
                        if j < len(hdr_cells):
                            hdr_cells[j].text = header
                            for run in hdr_cells[j].paragraphs[0].runs:
                                run.bold = True
                    for r_idx, row in enumerate(rows):
                        row_cells = table.rows[r_idx + 1].cells
                        for j, cell_text in enumerate(row):
                            if j < len(row_cells):
                                row_cells[j].text = cell_text
                continue

            # Listas numeradas
            elif re.match(r"^\d+\. ", line):
                text = re.sub(r"^\d+\. ", "", line)
                para = doc.add_paragraph(style="List Number")
                _add_inline_formatting(para, text)

            # Listas con viñetas
            elif line.startswith("- ") or line.startswith("* "):
                text = line[2:].strip()
                para = doc.add_paragraph(style="List Bullet")
                _add_inline_formatting(para, text)

            # Línea en blanco
            elif line.strip() == "":
                doc.add_paragraph("")

            # Párrafo normal con formato inline
            else:
                para = doc.add_paragraph()
                _add_inline_formatting(para, line)

            i += 1

        doc.save(output_path)
        print(f"[OK] Manual Word (.docx) guardado en: {output_path}")

    except ImportError:
        print("[ERROR] python-docx no está instalado. Instalar con: pip install python-docx")
        print("[INFO] Se guardará solo el archivo Markdown.")


def _add_inline_formatting(para, text: str) -> None:
    """Aplica formato inline (negrita, cursiva) a un párrafo de python-docx."""
    from docx.shared import RGBColor

    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|_\[[^\]]+\]_)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = para.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = para.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("_[") and part.endswith("]_"):
            run = para.add_run(part[2:-2] + " [pendiente]")
            run.italic = True
            run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        else:
            para.add_run(part)


def generate_manual(
    game_data: dict,
    template_path: str,
    output_dir: str,
    output_format: str = "both"
) -> dict:
    """
    Genera el manual completo en los formatos solicitados.
    Retorna un dict con las rutas de los archivos generados.
    """
    os.makedirs(output_dir, exist_ok=True)

    # Agregar fecha si no está en los datos
    if "FECHA" not in game_data or not game_data.get("FECHA"):
        game_data["FECHA"] = datetime.now().strftime("%Y-%m-%d")

    # Nombre base del archivo
    game_name = game_data.get("NOMBRE_DEL_JUEGO", "manual_juego")
    safe_name = "".join(c if c.isalnum() or c in "-_ " else "_" for c in game_name)
    safe_name = safe_name.replace(" ", "_").lower()

    content = render_template(template_path, game_data)
    generated_files = {}

    if output_format in ("md", "both"):
        md_path = os.path.join(output_dir, f"{safe_name}.md")
        save_markdown(content, md_path)
        generated_files["md"] = md_path

    if output_format in ("docx", "both"):
        docx_path = os.path.join(output_dir, f"{safe_name}.docx")
        save_docx(content, docx_path)
        generated_files["docx"] = docx_path

    return generated_files


def main():
    parser = argparse.ArgumentParser(description="Generador de Manuales de Juegos de Mesa")
    parser.add_argument("input_json", help="Ruta al archivo JSON con los datos del juego")
    parser.add_argument("--output-dir", default="./output", help="Directorio de salida")
    parser.add_argument(
        "--format",
        choices=["md", "docx", "both"],
        default="both",
        help="Formato de salida: md, docx, o both (ambos)"
    )
    parser.add_argument(
        "--template",
        default=None,
        help="Ruta al template Markdown personalizado"
    )

    args = parser.parse_args()

    # Buscar template por defecto
    if args.template is None:
        script_dir = Path(__file__).parent.parent
        default_template = script_dir / "templates" / "manual_template.md"
        if default_template.exists():
            args.template = str(default_template)
        else:
            print("[ERROR] No se encontró el template. Especifica uno con --template")
            sys.exit(1)

    game_data = load_game_data(args.input_json)
    files = generate_manual(game_data, args.template, args.output_dir, args.format)

    print("\n=== Archivos generados ===")
    for fmt, path in files.items():
        print(f"  {fmt.upper()}: {path}")


if __name__ == "__main__":
    main()
